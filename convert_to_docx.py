import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_word_document(md_path, docx_path):
    doc = Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Base styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_data = []

    def flush_table():
        nonlocal table_data
        if not table_data:
            return
        
        # Clean rows (remove markdown separator rows)
        cleaned_rows = [r for r in table_data if not re.match(r'^\s*\|?\s*[-:\s|]+\s*\|?\s*$', '|'.join(r))]
        if not cleaned_rows:
            table_data = []
            return

        cols_count = max(len(r) for r in cleaned_rows)
        tbl = doc.add_table(rows=len(cleaned_rows), cols=cols_count)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        for row_idx, row in enumerate(cleaned_rows):
            for col_idx in range(cols_count):
                cell = tbl.cell(row_idx, col_idx)
                text = row[col_idx] if col_idx < len(row) else ""
                cell.text = text.strip()
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                
                # Format cell text
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                
                if row_idx == 0:
                    set_cell_background(cell, "1E293B")
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        run.font.name = 'Calibri'
                        run.font.size = Pt(10)
                else:
                    bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                    set_cell_background(cell, bg)
                    for run in p.runs:
                        run.font.name = 'Calibri'
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        table_data = []

    def flush_code():
        nonlocal code_lines
        if not code_lines:
            return
        code_text = "\n".join(code_lines)
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.text = code_text
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        for run in p.runs:
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8) # Light blue
        
        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        code_lines = []

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')

        # Code block handling
        if line.strip().startswith('```'):
            if in_code_block:
                flush_code()
                in_code_block = False
            else:
                if in_table:
                    flush_table()
                    in_table = False
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Table handling
        if line.strip().startswith('|') and '|' in line.strip()[1:]:
            if not in_table:
                in_table = True
            # Parse columns
            cols = [c.strip() for c in line.strip().split('|')[1:-1]]
            if cols:
                table_data.append(cols)
            i += 1
            continue
        elif in_table:
            flush_table()
            in_table = False

        # Blank line
        if not line.strip():
            i += 1
            continue

        # Horizontal rule
        if line.strip() in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue

        # Images
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', line.strip())
        if img_match:
            caption, img_rel_path = img_match.groups()
            # resolve path
            if os.path.isabs(img_rel_path):
                img_full_path = img_rel_path
            else:
                img_full_path = os.path.join(os.path.dirname(md_path), img_rel_path)
            
            if os.path.exists(img_full_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(10)
                p_img.paragraph_format.space_after = Pt(4)
                run = p_img.add_run()
                run.add_picture(img_full_path, width=Inches(5.8))
                
                if caption:
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_after = Pt(12)
                    run_cap = p_cap.add_run(f"Figure: {caption}")
                    run_cap.font.italic = True
                    run_cap.font.size = Pt(9.5)
                    run_cap.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line[2:].strip())
            run.font.bold = True
            run.font.size = Pt(22)
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            run.font.name = 'Calibri'
        elif line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[3:].strip())
            run.font.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF) # Blue heading
            run.font.name = 'Calibri'
        elif line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[4:].strip())
            run.font.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            run.font.name = 'Calibri'
        elif line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line[5:].strip())
            run.font.bold = True
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
            run.font.name = 'Calibri'
        elif line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.25)
            clean_text = line[2:].strip()
            # replace formatting
            run = p.add_run(clean_text)
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        elif line.strip().startswith('- ') or line.strip().startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            bullet_text = line.strip()[2:].strip()
            # parse bold sections in bullet
            parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
        elif re.match(r'^\d+\.\s', line.strip()):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            num_text = re.sub(r'^\d+\.\s', '', line.strip())
            parts = re.split(r'(\*\*.*?\*\*)', num_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)
            parts = re.split(r'(\*\*.*?\*\*|`.*?`)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                elif part.startswith('`') and part.endswith('`'):
                    r = p.add_run(part[1:-1])
                    r.font.name = 'Consolas'
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
                else:
                    p.add_run(part)

        i += 1

    if in_table:
        flush_table()
    if in_code_block:
        flush_code()

    doc.save(docx_path)
    print(f"Successfully generated: {docx_path}")

if __name__ == '__main__':
    md_file = "/Users/aditisabharwal/arun practise/DOCUMENTATION.md"
    docx_file = "/Users/aditisabharwal/arun practise/Project_Tracker_Documentation.docx"
    create_word_document(md_file, docx_file)
